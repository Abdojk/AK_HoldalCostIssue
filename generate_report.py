#!/usr/bin/env python3
"""
D365 F&O — Unmarked Return Orders Moving Average Impact — Excel Report Generator
Author: Abdo J. Khoury | Info-Sys
"""

import sys
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ── Colour Constants ─────────────────────────────────────────────────────────
D365_NAVY       = "002060"
D365_MID_BLUE   = "1F6DB5"
D365_LIGHT_BLUE = "DEEAF1"
WHITE           = "FFFFFF"
BORDER_COLOUR   = "9DC3E6"
AUTHOR_BLUE     = "0070C0"
AMBER_BG        = "FCE4D6"
AMBER_TEXT      = "7F3F00"
FONT_DARK       = "1F2D3D"

FONT_NAME = "Aptos"

COLUMN_HEADERS = [
    "#", "Return Order Number", "Item Number", "Return Quantity",
    "Return Cost Applied (USD)", "Original Issue Cost (USD)", "Delta (USD)",
    "Return Financial Date", "Return Voucher", "Transaction Currency",
    "Receipt Status"
]

COLUMN_WIDTHS = [5, 22, 18, 16, 26, 26, 14, 22, 18, 22, 14]

THIN_BORDER = Border(
    left=Side(style="thin", color=BORDER_COLOUR),
    right=Side(style="thin", color=BORDER_COLOUR),
    top=Side(style="thin", color=BORDER_COLOUR),
    bottom=Side(style="thin", color=BORDER_COLOUR),
)


# ── Parsing ──────────────────────────────────────────────────────────────────

def safe_float(val):
    try:
        return float(val.strip().replace(",", ""))
    except (ValueError, AttributeError):
        return 0.00


def parse_date(val):
    try:
        return datetime.strptime(val.strip()[:10], "%Y-%m-%d").strftime("%d %b %Y")
    except (ValueError, AttributeError):
        return val.strip() if val else ""


def parse_data(raw_text):
    lines = raw_text.strip().splitlines()
    if not lines:
        return []
    rows = []
    for line in lines[1:]:  # skip header
        line = line.strip()
        if not line:
            continue
        cols = line.split("\t")
        if len(cols) < 11:
            cols.extend([""] * (11 - len(cols)))
        rows.append({
            "return_order":    cols[0].strip(),
            "item_number":     cols[1].strip(),
            "return_qty":      safe_float(cols[2]),
            "return_cost":     safe_float(cols[3]),
            "original_cost":   safe_float(cols[4]),
            "delta":           safe_float(cols[5]),
            "financial_date":  parse_date(cols[6]),
            "voucher":         cols[7].strip(),
            "currency":        cols[8].strip(),
            "receipt_status":  cols[9].strip(),
            "legal_entity":    cols[10].strip(),
        })
    return rows


# ── Classification ───────────────────────────────────────────────────────────

def classify_rows(rows):
    high_risk = []
    anomaly = []
    non_compliant = []

    for row in rows:
        # Rule 1 — High Risk
        if row["original_cost"] == 0.0 and row["return_cost"] > 0.0:
            high_risk.append(row)
        # Rule 2 — Anomaly
        elif row["return_cost"] == 0.0 and row["original_cost"] == 0.0:
            anomaly.append(row)
        # Rule 3 — Non-Compliant
        elif row["delta"] == 0.0 and row["original_cost"] > 0.0:
            non_compliant.append(row)
        # Unclassified → High Risk with note
        else:
            row["note"] = "UNCLASSIFIED — manual review required"
            high_risk.append(row)

    return high_risk, anomaly, non_compliant


# ── Styles ───────────────────────────────────────────────────────────────────

def header_font():
    return Font(name=FONT_NAME, bold=True, size=10, color=WHITE)

def header_fill():
    return PatternFill(start_color=D365_NAVY, end_color=D365_NAVY, fill_type="solid")

def header_alignment():
    return Alignment(horizontal="center", vertical="center", wrap_text=True)

def data_font():
    return Font(name=FONT_NAME, size=10, color=FONT_DARK)

def data_alignment():
    return Alignment(horizontal="center", vertical="center")

def even_fill():
    return PatternFill(start_color=D365_LIGHT_BLUE, end_color=D365_LIGHT_BLUE, fill_type="solid")

def odd_fill():
    return PatternFill(start_color=WHITE, end_color=WHITE, fill_type="solid")

def totals_fill():
    return PatternFill(start_color=D365_MID_BLUE, end_color=D365_MID_BLUE, fill_type="solid")

def totals_font():
    return Font(name=FONT_NAME, bold=True, size=10, color=WHITE)

def banner_font():
    return Font(name=FONT_NAME, bold=True, size=11, color=WHITE)

def author_font():
    return Font(name=FONT_NAME, italic=True, size=9, color=AUTHOR_BLUE)

def note_font():
    return Font(name=FONT_NAME, italic=True, size=10, color=AMBER_TEXT)

def note_fill():
    return PatternFill(start_color=AMBER_BG, end_color=AMBER_BG, fill_type="solid")


# ── Helpers ──────────────────────────────────────────────────────────────────

def set_column_widths(ws):
    for i, w in enumerate(COLUMN_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def write_header_row(ws, row_num):
    for col_idx, title in enumerate(COLUMN_HEADERS, 1):
        cell = ws.cell(row=row_num, column=col_idx, value=title)
        cell.font = header_font()
        cell.fill = header_fill()
        cell.alignment = header_alignment()
        cell.border = THIN_BORDER
    ws.row_dimensions[row_num].height = 30


def write_data_row(ws, row_num, seq, row_data):
    values = [
        seq,
        row_data["return_order"],
        row_data["item_number"],
        row_data["return_qty"],
        row_data["return_cost"],
        row_data["original_cost"],
        row_data["delta"],
        row_data["financial_date"],
        row_data["voucher"],
        row_data["currency"],
        row_data["receipt_status"],
    ]
    is_even = (seq % 2 == 0)
    fill = even_fill() if is_even else odd_fill()

    for col_idx, val in enumerate(values, 1):
        cell = ws.cell(row=row_num, column=col_idx, value=val)
        cell.font = data_font()
        cell.fill = fill
        cell.alignment = data_alignment()
        cell.border = THIN_BORDER
        # Number formats
        if col_idx == 4:  # Return Quantity — 3 decimals
            cell.number_format = "0.000"
        elif col_idx in (5, 6, 7):  # costs / delta
            cell.number_format = "#,##0.00"

    # If row has a note (unclassified), write it in col 12
    if "note" in row_data:
        cell = ws.cell(row=row_num, column=12, value=row_data["note"])
        cell.font = data_font()
        cell.fill = fill
        cell.alignment = data_alignment()
        cell.border = THIN_BORDER


def write_author_row(ws, row_num, entity, today_str):
    ws.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=11)
    cell = ws.cell(row=row_num, column=1)
    cell.value = (
        f"Legal Entity: {entity}  |  Prepared by: Abdo J. Khoury | Info-Sys  |  "
        f"Date: {today_str}  |  Source: D365 F&O INVENTTRANS — "
        "Unmarked Returns (MARKINGREFINVENTTRANSORIGIN = 0, SALESTYPE = 4)"
    )
    cell.font = author_font()
    cell.alignment = Alignment(horizontal="left", vertical="center")


# ── Excel Generation ─────────────────────────────────────────────────────────

def generate_excel(high_risk, anomaly, legal_entity, filename):
    entity = legal_entity.upper()
    today_str = datetime.now().strftime("%d %b %Y")
    wb = Workbook()

    # ── Tab 1: High Risk ─────────────────────────────────────────────────
    ws1 = wb.active
    ws1.title = "\U0001f534 High Risk"
    set_column_widths(ws1)

    write_header_row(ws1, 1)

    for i, row_data in enumerate(high_risk, 1):
        write_data_row(ws1, i + 1, i, row_data)

    last_data_row = len(high_risk) + 1
    totals_row = last_data_row + 1

    # Totals row
    ws1.cell(row=totals_row, column=1, value="TOTAL").font = totals_font()
    for col_idx in range(1, 12):
        cell = ws1.cell(row=totals_row, column=col_idx)
        cell.fill = totals_fill()
        cell.font = totals_font()
        cell.alignment = data_alignment()
        cell.border = THIN_BORDER

    if high_risk:
        # SUM formulas for Return Cost Applied (col 5) and Delta (col 7)
        ws1.cell(row=totals_row, column=5).value = f"=SUM(E2:E{last_data_row})"
        ws1.cell(row=totals_row, column=5).number_format = "#,##0.00"
        ws1.cell(row=totals_row, column=7).value = f"=SUM(G2:G{last_data_row})"
        ws1.cell(row=totals_row, column=7).number_format = "#,##0.00"

    # Author row
    author_row = totals_row + 2
    write_author_row(ws1, author_row, entity, today_str)

    ws1.freeze_panes = "B2"

    # ── Tab 2: Anomaly ───────────────────────────────────────────────────
    ws2 = wb.create_sheet(title="\u26AB Anomaly")
    set_column_widths(ws2)

    # Banner row
    ws2.merge_cells(start_row=1, start_column=1, end_row=1, end_column=11)
    banner_cell = ws2.cell(row=1, column=1)
    banner_cell.value = (
        f"{entity} — Return Orders: Zero Cost Anomaly | "
        "Requires Separate Investigation | Abdo J. Khoury | Info-Sys"
    )
    banner_cell.font = banner_font()
    banner_cell.fill = header_fill()
    banner_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws2.row_dimensions[1].height = 28

    # Header row
    write_header_row(ws2, 2)

    if anomaly:
        for i, row_data in enumerate(anomaly, 1):
            write_data_row(ws2, i + 2, i, row_data)

        last_data_row_2 = len(anomaly) + 2

        # Investigation note
        note_row = last_data_row_2 + 2
        ws2.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=11)
        note_cell = ws2.cell(row=note_row, column=1)
        note_cell.value = (
            "\u26A0  INVESTIGATION NOTE — Items returned at USD 0.00. Root cause: "
            "item had zero moving average cost at return date, indicating no prior "
            f"purchase receipt existed in {entity}. Investigate item purchase history "
            "and raise cost adjustment if required."
        )
        note_cell.font = note_font()
        note_cell.fill = note_fill()
        note_cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws2.row_dimensions[note_row].height = 55

        # Author row
        author_row_2 = note_row + 2
        write_author_row(ws2, author_row_2, entity, today_str)
    else:
        ws2.cell(row=3, column=1, value=f"No anomaly records found for {entity}.")
        ws2.cell(row=3, column=1).font = data_font()

    ws2.freeze_panes = "B3"

    wb.save(filename)


# ── Main ─────────────────────────────────────────────────────────────────────

def main(raw_text, legal_entity):
    rows = parse_data(raw_text)
    total_input = len(rows)
    high_risk, anomaly, non_compliant = classify_rows(rows)

    entity = legal_entity.upper()
    filename = f"{entity}_Unmarked_Returns_Moving_Average_Impact.xlsx"

    generate_excel(high_risk, anomaly, entity, filename)

    # Calculate total exposure
    total_exposure = sum(r["return_cost"] for r in high_risk)

    print()
    print("\u2550" * 48)
    print(f"  REPORT GENERATED — {entity}")
    print("\u2550" * 48)
    print(f"  File            : {filename}")
    print(f"  Total rows input: {total_input}")
    print(f"  \U0001f534 High Risk    : {len(high_risk)} rows  —  USD {total_exposure:,.2f} total exposure")
    print(f"  \u26AB Anomaly      : {len(anomaly)} rows")
    print(f"  \U0001f7e1 Non-compliant: {len(non_compliant)} rows (excluded from Excel — Delta = 0 by coincidence)")
    print("\u2550" * 48)
    print("  Prepared by: Abdo J. Khoury | Info-Sys")
    print()

    return filename


if __name__ == "__main__":
    print("Ready. Paste your query results below and tell me the legal entity code.")
    print("Format: tab-separated, header row included, copied directly from SSMS or Notepad.")
