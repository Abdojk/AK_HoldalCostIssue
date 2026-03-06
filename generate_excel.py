#!/usr/bin/env python3
"""
Unmarked Returns — Moving Average Impact Excel Generator
Author  : Abdo J. Khoury | Info-Sys
Purpose : Accepts tab-separated SSMS output from the Step 6 query and
          auto-generates a D365 F&O styled Excel report for any legal entity.

Usage:
    python generate_excel.py input_data.tsv
    python generate_excel.py input_data.tsv --output custom_name.xlsx
    cat input_data.tsv | python generate_excel.py -
"""

import argparse
import csv
import sys
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------------------
# D365 F&O Colour Scheme
# ---------------------------------------------------------------------------
NAVY = "002060"
MID_BLUE = "1F6DB5"
LIGHT_BLUE = "DEEAF1"
WHITE = "FFFFFF"
SOFT_BLUE = "9DC3E6"
DARK_NAVY = "1F2D3D"
D365_BLUE = "0070C0"
AMBER_BG = "FCE4D6"
DARK_AMBER = "7F3F00"

# ---------------------------------------------------------------------------
# Fonts
# ---------------------------------------------------------------------------
FONT_FAMILY = "Aptos"

HEADER_FONT = Font(name=FONT_FAMILY, bold=True, size=10, color=WHITE)
TOTALS_FONT = Font(name=FONT_FAMILY, bold=True, size=10, color=WHITE)
DATA_FONT = Font(name=FONT_FAMILY, size=10, color=DARK_NAVY)
AUTHOR_FONT = Font(name=FONT_FAMILY, italic=True, size=9, color=D365_BLUE)
ANOMALY_NOTE_FONT = Font(name=FONT_FAMILY, size=10, color=DARK_AMBER)
BANNER_FONT = Font(name=FONT_FAMILY, bold=True, size=12, color=WHITE)
EMPTY_MSG_FONT = Font(name=FONT_FAMILY, italic=True, size=10, color=DARK_NAVY)

# ---------------------------------------------------------------------------
# Fills
# ---------------------------------------------------------------------------
HEADER_FILL = PatternFill(start_color=NAVY, end_color=NAVY, fill_type="solid")
TOTALS_FILL = PatternFill(start_color=MID_BLUE, end_color=MID_BLUE, fill_type="solid")
EVEN_FILL = PatternFill(start_color=LIGHT_BLUE, end_color=LIGHT_BLUE, fill_type="solid")
ODD_FILL = PatternFill(start_color=WHITE, end_color=WHITE, fill_type="solid")
BANNER_FILL = PatternFill(start_color=NAVY, end_color=NAVY, fill_type="solid")
ANOMALY_NOTE_FILL = PatternFill(start_color=AMBER_BG, end_color=AMBER_BG, fill_type="solid")

# ---------------------------------------------------------------------------
# Border
# ---------------------------------------------------------------------------
THIN_BORDER = Border(
    left=Side(style="thin", color=SOFT_BLUE),
    right=Side(style="thin", color=SOFT_BLUE),
    top=Side(style="thin", color=SOFT_BLUE),
    bottom=Side(style="thin", color=SOFT_BLUE),
)

# ---------------------------------------------------------------------------
# Column definitions
# ---------------------------------------------------------------------------
COLUMNS = [
    ("#", 5),
    ("Return Order Number", 22),
    ("Item Number", 18),
    ("Return Quantity", 16),
    ("Return Cost Applied (USD)", 26),
    ("Original Issue Cost (USD)", 26),
    ("Delta (USD)", 14),
    ("Return Financial Date", 22),
    ("Return Voucher", 18),
    ("Transaction Currency", 22),
    ("Receipt Status", 14),
]

# Input column names as they come from the SSMS query
INPUT_COLS = [
    "Return Order Number",
    "Item Number",
    "Return Quantity",
    "Return Cost Applied",
    "Original Issue Cost",
    "Delta",
    "Return Financial Date",
    "Return Voucher",
    "Transaction Currency",
    "Receipt Status",
    "Legal Entity",
]


def parse_input(source):
    """Parse tab-separated input and return list of row dicts."""
    reader = csv.DictReader(source, delimiter="\t")
    # Strip whitespace from field names
    reader.fieldnames = [f.strip() for f in reader.fieldnames]
    rows = []
    for row in reader:
        cleaned = {k.strip(): v.strip() if v else "" for k, v in row.items()}
        rows.append(cleaned)
    return rows


def classify_rows(rows):
    """Split rows into high_risk, anomaly, and non_compliant lists."""
    high_risk = []
    anomaly = []
    non_compliant = []

    for row in rows:
        return_cost = float(row.get("Return Cost Applied", 0) or 0)
        orig_cost = float(row.get("Original Issue Cost", 0) or 0)
        delta = float(row.get("Delta", 0) or 0)

        if return_cost == 0.0 and orig_cost == 0.0:
            anomaly.append(row)
        elif orig_cost == 0.0 and return_cost > 0.0:
            high_risk.append(row)
        elif delta == 0.0 and orig_cost > 0.0:
            non_compliant.append(row)
        else:
            # Any row with a non-zero delta and non-zero orig cost is also
            # high risk (cost mismatch).
            high_risk.append(row)

    return high_risk, anomaly, non_compliant


def set_column_widths(ws):
    """Set column widths per specification."""
    for idx, (_, width) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width


def write_header_row(ws, row_num):
    """Write the styled header row."""
    for col_idx, (title, _) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=row_num, column=col_idx, value=title)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = THIN_BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")


def write_data_row(ws, row_num, data_row, row_number):
    """Write a single data row with alternating colours."""
    fill = EVEN_FILL if row_number % 2 == 0 else ODD_FILL

    values = [
        row_number,
        data_row.get("Return Order Number", ""),
        data_row.get("Item Number", ""),
        safe_float(data_row.get("Return Quantity", "")),
        safe_float(data_row.get("Return Cost Applied", "")),
        safe_float(data_row.get("Original Issue Cost", "")),
        safe_float(data_row.get("Delta", "")),
        data_row.get("Return Financial Date", ""),
        data_row.get("Return Voucher", ""),
        data_row.get("Transaction Currency", ""),
        data_row.get("Receipt Status", ""),
    ]

    for col_idx, val in enumerate(values, start=1):
        cell = ws.cell(row=row_num, column=col_idx, value=val)
        cell.font = DATA_FONT
        cell.fill = fill
        cell.border = THIN_BORDER
        if col_idx == 1:
            cell.alignment = Alignment(horizontal="center")
        elif col_idx in (4, 5, 6, 7):
            cell.number_format = '#,##0.00'
            cell.alignment = Alignment(horizontal="right")


def safe_float(val):
    """Convert to float if possible, otherwise return the original string."""
    if val == "" or val is None:
        return 0.0
    try:
        return float(str(val).replace(",", ""))
    except ValueError:
        return val


def write_totals_row(ws, row_num, data_start_row, data_end_row):
    """Write the totals row with SUM formulas (not hardcoded)."""
    # Column E = Return Cost Applied, Column G = Delta
    cost_col = get_column_letter(5)  # E
    delta_col = get_column_letter(7)  # G

    labels_and_formulas = {
        1: "TOTALS",
        5: f"=SUM({cost_col}{data_start_row}:{cost_col}{data_end_row})",
        7: f"=SUM({delta_col}{data_start_row}:{delta_col}{data_end_row})",
    }

    for col_idx in range(1, len(COLUMNS) + 1):
        cell = ws.cell(row=row_num, column=col_idx)
        if col_idx in labels_and_formulas:
            cell.value = labels_and_formulas[col_idx]
        cell.font = TOTALS_FONT
        cell.fill = TOTALS_FILL
        cell.border = THIN_BORDER
        if col_idx in (5, 7):
            cell.number_format = '#,##0.00'
            cell.alignment = Alignment(horizontal="right")
        elif col_idx == 1:
            cell.alignment = Alignment(horizontal="center")


def write_author_line(ws, row_num):
    """Write the author footer line."""
    today = date.today().strftime("%d %B %Y")
    author_text = f"Abdo J. Khoury | Info-Sys | {today}"
    cell = ws.cell(row=row_num, column=2, value=author_text)
    cell.font = AUTHOR_FONT
    cell.alignment = Alignment(horizontal="left")


def build_high_risk_tab(wb, high_risk_rows):
    """Build Tab 1 — High Risk."""
    ws = wb.active
    ws.title = "\U0001f534 High Risk"

    set_column_widths(ws)
    write_header_row(ws, 1)

    # Data rows
    for idx, row in enumerate(high_risk_rows, start=1):
        write_data_row(ws, idx + 1, row, idx)

    data_start = 2
    data_end = len(high_risk_rows) + 1

    # Totals row
    totals_row = data_end + 1
    if high_risk_rows:
        write_totals_row(ws, totals_row, data_start, data_end)
    else:
        totals_row = data_end  # no data, skip totals

    # Author line
    write_author_line(ws, totals_row + 2)

    # Freeze panes at B2
    ws.freeze_panes = "B2"


def build_anomaly_tab(wb, anomaly_rows, legal_entity):
    """Build Tab 2 — Anomaly."""
    ws = wb.create_sheet(title="\u26ab Anomaly")

    set_column_widths(ws)

    # Row 1: Navy banner
    for col_idx in range(1, len(COLUMNS) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = BANNER_FILL
        cell.border = THIN_BORDER
    banner_cell = ws.cell(row=1, column=2, value="ANOMALY RECORDS")
    banner_cell.font = BANNER_FONT
    banner_cell.fill = BANNER_FILL

    # Row 2: Header
    write_header_row(ws, 2)

    if anomaly_rows:
        for idx, row in enumerate(anomaly_rows, start=1):
            write_data_row(ws, idx + 2, row, idx)
        last_data_row = len(anomaly_rows) + 2
    else:
        # Empty message
        empty_cell = ws.cell(
            row=3,
            column=2,
            value=f"No anomaly records found for {legal_entity}.",
        )
        empty_cell.font = EMPTY_MSG_FONT
        last_data_row = 3

    # Amber investigation note
    note_row = last_data_row + 2
    note_text = (
        "Investigation required: These records have zero cost on both the "
        "return receipt and the original issue transaction. Review inventory "
        "transactions and costing entries for potential data integrity issues."
    )
    for col_idx in range(1, len(COLUMNS) + 1):
        cell = ws.cell(row=note_row, column=col_idx)
        cell.fill = ANOMALY_NOTE_FILL
        cell.border = THIN_BORDER
    note_cell = ws.cell(row=note_row, column=2, value=note_text)
    note_cell.font = ANOMALY_NOTE_FONT
    note_cell.fill = ANOMALY_NOTE_FILL
    # Merge note across columns for readability
    ws.merge_cells(
        start_row=note_row, start_column=2,
        end_row=note_row, end_column=len(COLUMNS),
    )

    # Author line
    write_author_line(ws, note_row + 2)

    # Freeze panes at B3
    ws.freeze_panes = "B3"


def main():
    parser = argparse.ArgumentParser(
        description="Generate D365 F&O styled Excel from unmarked returns query output."
    )
    parser.add_argument(
        "input",
        help="Path to tab-separated input file, or '-' for stdin.",
    )
    parser.add_argument(
        "--output", "-o",
        help="Output Excel file name (default: auto-generated from legal entity).",
        default=None,
    )
    args = parser.parse_args()

    # Read input
    if args.input == "-":
        rows = parse_input(sys.stdin)
    else:
        with open(args.input, "r", encoding="utf-8-sig") as f:
            rows = parse_input(f)

    if not rows:
        print("ERROR: No data rows found in input.")
        sys.exit(1)

    # Determine legal entity from first row
    legal_entity = rows[0].get("Legal Entity", "UNKNOWN").strip()

    # Classify
    high_risk, anomaly, non_compliant = classify_rows(rows)

    # Build workbook
    wb = Workbook()
    build_high_risk_tab(wb, high_risk)
    build_anomaly_tab(wb, anomaly, legal_entity)

    # Output filename
    output_file = args.output or f"{legal_entity}_Unmarked_Returns_Moving_Average_Impact.xlsx"
    wb.save(output_file)

    # Summary to stdout
    total = len(rows)
    hr_total = sum(
        safe_float(r.get("Delta", 0)) for r in high_risk
        if isinstance(safe_float(r.get("Delta", 0)), float)
    )
    print(f"{'=' * 56}")
    print(f"EXCEL GENERATED — {legal_entity}")
    print(f"{'=' * 56}")
    print(f"File delivered      : {output_file}")
    print(f"Total rows analysed : {total}")
    print(f"\U0001f534 High Risk        : {len(high_risk)} lines — USD {hr_total:,.2f} total exposure")
    print(f"\u26ab Anomaly          : {len(anomaly)} lines")
    print(f"\U0001f7e1 Non-compliant    : {len(non_compliant)} lines (excluded from Excel)")
    print(f"{'=' * 56}")


if __name__ == "__main__":
    main()
